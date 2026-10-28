"""Build a direct-file index and a short reservation demonstration guide."""
from urllib.parse import quote

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt

from document_common import DOC_ROOT, FACTS, OUTPUTS, new_document, page, paragraph, save_document, table


def file_link(document, key, label):
    """Add a portable relative link to a current Word or PowerPoint file."""
    target=OUTPUTS[key].relative_to(DOC_ROOT).as_posix()
    relation=document.part.relate_to(quote(target),RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
    item=document.add_paragraph()
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),relation)
    run=OxmlElement('w:r')
    properties=OxmlElement('w:rPr')
    underline=OxmlElement('w:u')
    underline.set(qn('w:val'),'single')
    properties.append(underline)
    run.append(properties)
    value=OxmlElement('w:t')
    value.text=label
    run.append(value)
    link.append(run)
    item._p.append(link)


def build():
    """Create a two-page index that opens each deliverable without a ZIP."""
    doc=new_document('Навигация по комплекту')
    doc.sections[0].different_first_page_header_footer=False
    title=doc.add_paragraph('Материалы проекта Молочная ферма Луговая',style='Title')
    for run in title.runs:
        run.font.size=Pt(24)
    paragraph(doc,'Молочная ферма Луговая. Версия '+FACTS['configuration_version']+'. Автор: '+FACTS['author']+'.')
    paragraph(doc,'Откройте нужный документ по ссылке ниже. Актуальные материалы доступны отдельными файлами в папке Документация.')
    doc.add_heading('Результат доработки',level=2)
    table(doc,['Состояние молока','Склад','Резерв','Свободно'],[
        ('До брони','75 л','0 л','75 л'),('Бронь 20 л','75 л','20 л','55 л'),
        ('Продажа по брони','55 л','0 л','55 л'),('Отмена до продажи','75 л','0 л','75 л'),
    ],[3.2,1.15,1.15,1.15])
    paragraph(doc,'Продажа и отмена — альтернативные ветви. Обычная продажа 56 л при свободных 55 л отклоняется.')
    doc.add_heading('Открыть материалы',level=2)
    for key,label,description in [
        ('presentation','Презентация PowerPoint','Обзор проекта на 11 слайдах: назначение, структура, интерфейс, учёт, бронирование и проверки. Заметки на выступление до пяти минут.'),
        ('general','Общая документация Word','Обзор изменений на странице аннотации; структура и алгоритмы, подробная модель брони в разделе 9.'),
        ('user','Инструкция пользователя Word','Бронирование и проверка результатов находятся в разделах 2–3, сразу после входа.'),
        ('testing','Документация по тестированию Word','Smoke-тест в разделе 3; подробные сценарии по ролям в разделах 4–6; журнал и дефекты в разделе 8. Автоматические результаты в приложении А.'),
        ('design','Обоснование дизайна Word','Оформление новых команд, состояния и отчёта рассматривается в первом разделе.'),
    ]:
        file_link(doc,key,label)
        paragraph(doc,description)
    page(doc,'Короткая демонстрация на занятии')
    for index,text in enumerate([
        'Открыть отчёт исходного примера: молоко на складе 75 л, резерв 0, свободно 75 л.',
        'В Продаже создать и провести бронь 20 л. В отчёте проверить 75, 20 и 55.',
        'Попробовать обычную продажу 56 л. Показать отказ и сохранение остатков.',
        'Отменить бронь. Показать, что все 75 л снова свободны.',
        'Снова провести бронь и нажать «Продать по брони». Показать связанный Расход и остатки 55, 0 и 55.',
    ],1):
        paragraph(doc,f'{index}. {text}')
    paragraph(doc,'Демонстрация выполняется в отдельной копии базы. Чтобы вернуть исходное состояние после продажи, сначала отмените проведение Расхода, затем отмените бронь.')
    doc.add_heading('Что подтверждено',level=2)
    paragraph(doc,'В шести протоколах — 83 успешные автоматические проверки. Файлы лежат в 04_Тестирование/Протоколы. Сохранённые DT и CF находятся в deliverables; рабочая база — demo_base_v1.2.0. Запуск: Запустить_ферму_v1.2.0.command.')
    paragraph(doc,'Бронь продаётся целиком; автоматического срока отмены нет. Учебная лицензия запрещает печать и сохранение табличного документа. Щелчки мышью и читабельность следует проверить вручную. Старые снимки оставлены в приложениях с указанием, что они сделаны до доработки.')
    doc.add_heading('Вопросы для собрания',level=2)
    for text in ['Что я делал вчера, чтобы помочь проекту?', 'Что я буду делать сегодня, чтобы помочь проекту?', 'Какие проблемы мешают успешной работе над проектом?']:
        paragraph(doc,text)
    paragraph(doc,'Это анонс занятия. Ответы даются по фактической работе; отдельная подготовка к моделированию собрания не требуется.')
    save_document(doc,'navigation','Бронирование продукции Навигация по комплекту')


if __name__=='__main__':
    build()
