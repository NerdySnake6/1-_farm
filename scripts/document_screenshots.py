"""Add supplied screenshots to the project documentation."""
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def add_cropped_screenshot(document, filename, size, bounds, description, width=6.9):
    """Embed an original screenshot with an editable Office crop."""
    source_width, source_height = size
    left, top, right, bottom = bounds
    height = width * (bottom - top) / (right - left)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(
        str(ROOT / 'assets/screenshots' / filename),
        width=Inches(width), height=Inches(height),
    )
    picture._inline.docPr.set('descr', description)
    fill = picture._inline.xpath('.//pic:blipFill')[0]
    crop = OxmlElement('a:srcRect')
    for edge, value in {
        'l': left / source_width,
        't': top / source_height,
        'r': (source_width - right) / source_width,
        'b': (source_height - bottom) / source_height,
    }.items():
        crop.set(edge, str(round(value * 100000)))
    fill.insert(1, crop)


def add_screenshot_appendix(document):
    """Append a page showing the home form and the report result."""
    title = document.add_heading('Приложение А Экранные формы системы', level=1)
    title.paragraph_format.page_break_before = True
    document.add_paragraph(
        'На рисунках приведены снимки работающей конфигурации. '
        'Для печати показаны отдельные области окон. Содержимое форм и значения данных сохранены.'
    )
    document.add_heading('Рабочий стол', level=2)
    add_cropped_screenshot(
        document, 'Рабочий_стол.png', (1624, 1252), (35, 475, 1259, 995),
        'Рабочий стол фермы с кнопками документов, справочников и отчетов',
    )
    caption = document.add_paragraph('Рисунок А.1 Рабочий стол фермы')
    caption.paragraph_format.space_after = Pt(8)
    document.add_paragraph(
        'С рабочего стола доступны создание документов, журналы поступлений и продаж, '
        'номенклатура с фотографиями, цены и отчёты. Зелёным выделены основные действия. '
        'Кнопки открытия журналов имеют светлый фон.'
    )
    document.add_heading('Результат отчёта по остаткам', level=2)
    add_cropped_screenshot(
        document, 'Остатки.png', (1906, 840), (875, 530, 1752, 680),
        'Фрагмент отчета: молоко 75, творог 15, сметана 12',
    )
    document.add_paragraph('Рисунок А.2 Остатки продукции в демонстрационной базе')
    document.add_paragraph(
        'Фрагмент отчёта показывает наименование продукции и её количество. '
        'Остатки контрольного примера составляют 75 л молока, 15 кг творога и 12 кг сметаны. '
        'Данные соответствуют расчётам, приведённым в разделе испытаний.'
    )
