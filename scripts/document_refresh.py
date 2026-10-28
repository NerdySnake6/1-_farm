"""Prioritize booking workflows and visible changes in the Word document set."""
from docx.oxml.ns import qn

from document_common import page, paragraph, table


def reorder_sections(document, ordered_titles, renamed):
    """Move complete chapter blocks and rename headings before creating contents."""
    body = document.element.body
    children = list(body)
    headings = {p._p: p for p in document.paragraphs if p.style.name == 'Heading 1'}
    starts = [(i, headings[node]) for i, node in enumerate(children) if node in headings]
    blocks = {}
    for position, (start, heading) in enumerate(starts):
        end = starts[position + 1][0] if position + 1 < len(starts) else len(children)
        blocks[heading.text] = [node for node in children[start:end] if node.tag != qn('w:sectPr')]
    assert set(blocks) == set(ordered_titles), (set(blocks), set(ordered_titles))
    for nodes in blocks.values():
        for node in nodes:
            body.remove(node)
    section_properties = body.find(qn('w:sectPr'))
    for title in ordered_titles:
        nodes = blocks[title]
        headings[nodes[0]].text = renamed.get(title, title)
        for node in nodes:
            if section_properties is None:
                body.append(node)
            else:
                section_properties.addprevious(node)


def user(document):
    """Move booking to chapter two and add a separate page of result checks."""
    page(document, 'Проверка результатов бронирования')
    paragraph(document, 'Для одной и той же брони сравниваются два альтернативных исхода. Начальный пример: молоко на складе 75 л, действующих резервов нет. Количество сверяется в отчёте для той же организации и склада.')
    table(document, ['Состояние', 'На складе', 'В резерве', 'Свободно'], [
        ('До брони', '75 л', '0 л', '75 л'),
        ('После брони 20 л', '75 л', '20 л', '55 л'),
        ('Если продать по брони', '55 л', '0 л', '55 л'),
        ('Если отменить бронь до продажи', '75 л', '0 л', '75 л'),
    ], [3.2, 1.15, 1.15, 1.15])
    document.add_heading('Как отличить бронь от продажи', level=2)
    paragraph(document, 'После бронирования складское количество остаётся 75 л. После продажи оно уменьшается до 55 л. Нулевой резерв сам по себе не означает продажу: проверьте складское количество и связанный документ Расход.')
    document.add_heading('Проверка запрета', level=2)
    paragraph(document, 'При действующей брони 20 л попробуйте оформить обычную продажу 56 л. Система должна отказать, а в отчёте должны остаться значения 75, 20 и 55. Не уменьшайте бронь и не создавайте дополнительный приход для обхода нехватки.')
    document.add_heading('Возврат к исходному состоянию', level=2)
    paragraph(document, 'Если продажа по брони уже выполнена, отмените проведение её Расхода: склад станет 75, резерв 20, свободно 55. Затем отмените саму бронь: все 75 л снова будут свободны. Для повторного показа проверьте, что других активных броней нет.')
    paragraph(document, 'Сохранение и печать табличного документа заблокированы установленной учебной лицензией. Результат проверяйте в окне отчёта. Это ограничение платформы зафиксировано отдельно в протоколе platform-limitations.txt.')
    order = ['1 Начало работы и доступ', '8 Бронирование продукции', 'Проверка результатов бронирования', '2 Справочники и фотографии', '3 Установка и изменение цен', '4 Поступление продукции', '5 Продажа и исправление документов', '6 Отчёты и сверка результатов', '7 Ошибки и завершение работы']
    rename = dict(zip(order, ['1 Начало работы и доступ', '2 Бронирование продукции', '3 Проверка результатов бронирования', '4 Справочники и фотографии', '5 Установка и изменение цен', '6 Поступление продукции', '7 Продажа и исправление документов', '8 Отчёты и сверка результатов', '9 Ошибки и завершение работы']))
    for item in document.paragraphs:
        if item.style.name != 'Heading 1':
            updated = item.text.replace('из раздела 4', 'из раздела 6').replace('разделов 4 и 5', 'разделов 6 и 7').replace('по инструкции раздела 6', 'по инструкции раздела 8')
            if updated != item.text:
                item.text = updated
    reorder_sections(document, order, rename)


def design(document):
    """Lead the design explanation with the new booking interface decisions."""
    order = ['6 Оформление бронирования', '1 Назначение оформления', '2 Цвета и типографика', '3 Размещение элементов и навигация', '4 Изображения и оформление отчётов', '5 Проверка оформления']
    renamed = dict(zip(order, ['1 Оформление бронирования', '2 Назначение оформления', '3 Цвета и типографика', '4 Размещение элементов и навигация', '5 Изображения и оформление отчётов', '6 Проверка оформления']))
    reorder_sections(document, order, renamed)


def testing(document):
    """Place booking acceptance scenarios before the regression test details."""
    order = ['1 Объект и условия тестирования', '6 Бронирование: создание и отмена', '7 Бронирование: продажа и восстановление', '8 Команды брони и приёмка доработки', '2 Проверка учёта', '2 Проверка учёта (продолжение)', '3 Проверка форм и команд', '4 Проверка прав пользователей', '5 Результаты и оставшиеся проверки', 'Приложение А Снимки конфигурации']
    names = ['1 Объект и условия тестирования', '2 Создание и отмена брони', '3 Продажа по брони и восстановление резерва', '4 Команды брони и приёмка доработки', '5 Проверка прежнего учёта', '5 Проверка прежнего учёта (продолжение)', '6 Прежние формы и команды', '7 Права пользователей', '8 Результаты и ограничения', 'Приложение А Снимки конфигурации']
    for item in document.paragraphs:
        if item.style.name != 'Heading 1' and 'результаты бронирования приведены в разделах 6–8' in item.text:
            item.text = item.text.replace('разделах 6–8', 'разделах 2–4')
    reorder_sections(document, order, dict(zip(order, names)))
