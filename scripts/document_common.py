"""Shared facts, paths and Word formatting for the documentation set."""
from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_ROOT = ROOT / 'Документация'
PROTOCOLS = DOC_ROOT / '04_Тестирование/Протоколы'
FACTS = json.loads((ROOT/'scripts/documentation_facts.json').read_text())
AUTHOR = FACTS['author']
OUTPUTS = {
    'navigation': DOC_ROOT/'00_Навигация_по_комплекту.docx',
    'general': DOC_ROOT/'01_Общее_описание/Общая_документация.docx',
    'design': DOC_ROOT/'02_Дизайн/Обоснование_выбора_дизайна.docx',
    'user': DOC_ROOT/'03_Пользователю/Инструкция_пользователя.docx',
    'testing': DOC_ROOT/'04_Тестирование/Документация_по_тестированию.docx',
    'presentation': DOC_ROOT/'05_Презентация/Презентация_проекта.pptx',
}


def read_results(filename):
    """Read successful checks without treating error protocols as passing."""
    text = (PROTOCOLS/filename).read_text(encoding='utf-8-sig')
    assert 'ERROR' not in text and 'FAIL' not in text, filename
    return [line.split(' | ',1)[1] for line in text.splitlines() if line.startswith('PASS | ')]


def paragraph(document, text, bold=False):
    """Append a body paragraph using the shared typography."""
    item = document.add_paragraph()
    item.add_run(text).bold = bold
    return item


def page(document, title):
    """Begin a section on a new page."""
    item = document.add_heading(title, level=1)
    item.paragraph_format.page_break_before = True
    return item


def table(document, headers, rows, widths):
    """Create a bordered table with repeating headings and flexible rows."""
    item = document.add_table(rows=1, cols=len(headers))
    item.autofit = False
    for column, width in zip(item.columns, widths):
        column.width = Inches(width)
    for cell, text in zip(item.rows[0].cells, headers):
        cell.text = text
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'),'E8E8E8')
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    item.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values in rows:
        for cell, value in zip(item.add_row().cells, values):
            cell.text = str(value)
    for row in item.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for index, (cell, width) in enumerate(zip(row.cells,widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                border = OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    border.set(qn('w:'+key),value)
                borders.append(border)
            properties.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                margin = OxmlElement('w:'+edge)
                margin.set(qn('w:w'),'85')
                margin.set(qn('w:type'),'dxa')
                margins.append(margin)
            properties.append(margins)
            for body in cell.paragraphs:
                body.paragraph_format.space_after = Pt(3)
                body.paragraph_format.space_before = Pt(3)
                body.paragraph_format.line_spacing = 1.05
                if widths[index] <= 1.55:
                    body.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in body.runs:
                    run.font.size = Pt(10.5)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return item


def new_document(label):
    """Create a consistent A4 report with a blank first-page header."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2)
    section.different_first_page_header_footer = True
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0,0,0)
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 2'].font.size = Pt(13)
    section.header.paragraphs[0].text = 'МОЛОЧНАЯ ФЕРМА ЛУГОВАЯ   /   '+label.upper()
    for run in section.header.paragraphs[0].runs:
        run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
    return doc


def cover(document, title):
    """Add the agreed title page with the supplied author only."""
    for text,size,before,bold,right,is_title in [
        (title.upper(),18,125,True,False,False),
        ('Молочная ферма Луговая',24,10,True,False,True),
        ('Учёт продукции и бронирование в 1С версии 1.2.0',12,8,False,False,False),
        ('Дисциплина «Управление проектами»',12,0,False,False,False),
        ('Выполнил',12,80,False,True,False),
        (AUTHOR,12,0,True,True,False),
        ('2026',12,155,False,False,False),
    ]:
        item = document.add_paragraph(style='Title' if is_title else 'Normal')
        item.alignment = WD_ALIGN_PARAGRAPH.LEFT if right else WD_ALIGN_PARAGRAPH.CENTER
        item.paragraph_format.space_before = Pt(before)
        item.paragraph_format.space_after = Pt(10)
        if right:
            item.paragraph_format.left_indent = Cm(8)
        run = item.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold


def add_contents(document, entries):
    """Insert linked contents after the cover and bookmark each chapter."""
    first = next(p for p in document.paragraphs if p.style.name=='Heading 1')
    entry_paragraphs = []
    title = document.add_heading('Содержание', level=1)
    title.paragraph_format.page_break_before = True
    first._p.addprevious(title._p)
    for index, text in enumerate(entries):
        target = next(p for p in document.paragraphs if p.style.name=='Heading 1' and p.text==text)
        anchor = f'chapter_{index+1}'
        bookmark = OxmlElement('w:bookmarkStart')
        bookmark.set(qn('w:id'),str(index+1000))
        bookmark.set(qn('w:name'),anchor)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'),str(index+1000))
        target._p.insert(0,bookmark)
        target._p.append(end)
        item = document.add_paragraph()
        item.paragraph_format.space_after = Pt(9)
        item.paragraph_format.tab_stops.add_tab_stop(Cm(16.65),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        link = OxmlElement('w:hyperlink')
        link.set(qn('w:anchor'),anchor)
        run = OxmlElement('w:r')
        value = OxmlElement('w:t')
        value.text = text
        run.append(value)
        link.append(run)
        item._p.append(link)
        item.add_run('\t')
        item.add_run('…')
        first._p.addprevious(item._p)
        entry_paragraphs.append(item)
    return entry_paragraphs


def save_document(document, key, title):
    """Save a Word artifact with common metadata and sanitized titles."""
    document.core_properties.author = AUTHOR
    document.core_properties.last_modified_by = AUTHOR
    document.core_properties.title = title
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for border in document.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    OUTPUTS[key].parent.mkdir(parents=True,exist_ok=True)
    document.save(OUTPUTS[key])
    print(OUTPUTS[key])
