"""Build, render and synchronize all current Word documents."""
from pathlib import Path
import argparse
import importlib
import subprocess
import sys

import pdfplumber
from docx import Document
from docx.oxml.ns import qn

from document_common import OUTPUTS, ROOT

RENDERER = Path('/Users/nerdysnake6/.codex/plugins/cache/openai-primary-runtime/documents/26.1020.11127/skills/documents/render_docx.py')
BUILDERS = ['build_general_documentation','build_design_documentation','build_user_guide','build_testing_documentation','build_documentation_navigation']


def synchronize_contents(path, pdf_path):
    """Resolve chapter page numbers from the rendered artifact."""
    doc=Document(path)
    with pdfplumber.open(pdf_path) as pdf:
        page_lines=[page.extract_text().splitlines() for page in pdf.pages]
    targets={}
    for item in doc.paragraphs:
        for bookmark in item._p.findall(qn('w:bookmarkStart')):
            targets[bookmark.get(qn('w:name'))]=item.text
    count=0
    for item in doc.paragraphs:
        links=item._p.findall(qn('w:hyperlink'))
        if not links:
            continue
        anchor=links[0].get(qn('w:anchor'))
        if not anchor or anchor not in targets:
            continue
        heading=targets[anchor]
        pages=[i+1 for i,lines in enumerate(page_lines) if heading in lines]
        if len(pages)!=1:
            raise ValueError(f'Ambiguous heading {heading}: {pages}')
        assert item.runs
        item.runs[-1].text=('\t' if len(item.runs)==1 else '')+str(pages[0])
        count+=1
    if count:
        doc.save(path)
    return count


def render_and_sync(path):
    """Render a document, resolve contents and render its final layout."""
    folder=ROOT/'.build/documentation/qa'/path.parent.name
    folder.mkdir(parents=True,exist_ok=True)
    for stale in [*folder.glob('page-*.png'),*folder.glob('*.pdf')]:
        stale.unlink()
    command=[sys.executable,str(RENDERER),str(path),'--output_dir',str(folder),'--emit_pdf']
    subprocess.run(command,check=True)
    pdf=next(folder.glob('*.pdf'))
    if synchronize_contents(path,pdf):
        subprocess.run(command,check=True)
    return folder


def main():
    """Run the document build and optional render verification."""
    parser=argparse.ArgumentParser()
    parser.add_argument('--render',action='store_true')
    args=parser.parse_args()
    for name in BUILDERS:
        importlib.import_module(name).build()
    if args.render:
        for key,path in OUTPUTS.items():
            if key!='presentation':
                render_and_sync(path)


if __name__=='__main__':
    main()
