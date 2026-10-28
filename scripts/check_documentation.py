"""Validate documentation facts, navigation and rendered contents without 1C."""
import hashlib
import json
import re
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from document_common import AUTHOR, DOC_ROOT, FACTS, OUTPUTS, ROOT, read_results


def digest(path):
    """Return a file's SHA-256 digest."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_word(path):
    """Validate Word bookmarks, local links and cached contents page numbers."""
    document = Document(path)
    assert document.core_properties.author == AUTHOR, path
    targets = {}
    for paragraph in document.paragraphs:
        for bookmark in paragraph._p.findall(qn('w:bookmarkStart')):
            targets[bookmark.get(qn('w:name'))] = paragraph.text
    pdf_files = list((ROOT / '.build/documentation/qa' / path.parent.name).glob('*.pdf'))
    assert len(pdf_files) == 1, f'Render required: {path}'
    with pdfplumber.open(pdf_files[0]) as pdf:
        page_lines = [(page.extract_text() or '').splitlines() for page in pdf.pages]
    assert all(len(lines) > 2 for lines in page_lines), f'Empty page: {path}'
    for paragraph in document.paragraphs:
        for link in paragraph._p.findall(qn('w:hyperlink')):
            anchor = link.get(qn('w:anchor'))
            if not anchor:
                continue
            assert anchor in targets, (path, anchor)
            pages = [i + 1 for i, lines in enumerate(page_lines) if targets[anchor] in lines]
            assert len(pages) == 1, (path, targets[anchor], pages)
            assert int(paragraph.runs[-1].text.strip()) == pages[0], (path, anchor)
    links = 0
    for relation in document.part.rels.values():
        if relation.reltype.endswith('/hyperlink') and relation.is_external:
            target = urlsplit(relation.target_ref)
            if target.scheme:
                continue
            assert (path.parent / unquote(target.path)).exists(), (path, target.path)
            links += 1
    return {'pages': len(page_lines), 'local_links': links, 'contents_entries': len(targets)}


def check_markdown(path, pending_package=False):
    """Check portable local Markdown links in the current technical material."""
    for value in re.findall(r'\]\(([^)]+)\)', path.read_text()):
        target = urlsplit(value.strip('<>'))
        if not target.scheme and target.path:
            resolved = (path.parent / unquote(target.path)).resolve()
            bundle = ROOT / f"deliverables/Комплект_Молочная_ферма_v{FACTS['configuration_version']}.zip"
            if not (pending_package and resolved == bundle):
                assert resolved.exists(), (path, value)


def check(pending_package=False):
    """Check the current set, source evidence and optional preservation records."""
    active = {p for p in DOC_ROOT.rglob('*') if p.suffix in {'.docx', '.pptx'} and 'Архив' not in p.parts}
    assert active == set(OUTPUTS.values()), active.symmetric_difference(OUTPUTS.values())
    results = {key: check_word(path) for key, path in OUTPUTS.items() if path.suffix == '.docx'}
    assert results['navigation']['local_links'] == 5
    for path in [ROOT / 'README.md', *(DOC_ROOT / '06_Технические_материалы').glob('*.md')]:
        check_markdown(path, pending_package=pending_package)
    names = ['runtime-tests.txt', 'form-tests.txt', 'role-Покупка.txt', 'role-Продажа.txt', 'reservation-tests.txt', 'reservation-form-tests.txt']
    counts = [len(read_results(name)) for name in names]
    assert counts == list(FACTS['checks'].values()) == [21, 10, 9, 9, 28, 6]
    example = FACTS['control_example']
    assert [a - b for a, b in zip(example['receipt'], example['sale'])] == example['stock']
    assert sum(a * b for a, b in zip(example['sale'], example['prices'])) == example['sales_total'] == 4890
    source = etree.parse(str(ROOT / 'configuration/Configuration.xml'))
    assert source.xpath('//*[local-name()="Version"]/text()') == [FACTS['configuration_version']]
    with ZipFile(OUTPUTS['presentation']) as package:
        slides = [name for name in package.namelist() if re.fullmatch(r'ppt/slides/slide\d+.xml', name)]
        notes = [name for name in package.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+.xml', name)]
        assert len(slides) == len(notes) and len(slides) > 0
        assert package.testzip() is None
    archive_manifest = DOC_ROOT / 'Архив/2026-10-23/Реестр_архива.json'
    if archive_manifest.exists():
        for item in json.loads(archive_manifest.read_text()):
            assert digest(ROOT / item['archive']) == item['sha256'], item['archive']
    protected = ROOT / '.build/reservation/protected-bases.json'
    if protected.exists():
        for name, checksum in json.loads(protected.read_text()).items():
            assert digest(ROOT / name) == checksum, f'Protected file changed: {name}'
    results['checks'] = sum(counts)
    results['slides'] = len(slides)
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return results


if __name__ == '__main__':
    check()
